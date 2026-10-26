import re
import uuid
import shutil
import traceback
from datetime import datetime
from pathlib import Path

from flask import (
    Blueprint,
    render_template,
    request,
    redirect,
    url_for,
    flash,
    session,
    jsonify,
    send_file,
    current_app
)

from flask_login import login_required, current_user
from sqlalchemy import select
from werkzeug.utils import secure_filename

from database import db
from models import (
    Student,
    Subject,
    AttendanceSession,
    AttendanceRecord
)


attendance_bp = Blueprint(
    "attendance",
    __name__,
    url_prefix="/attendance"
)

ALLOWED_EXTENSIONS = {".docx", ".xlsx"}


def upload_folder():
    folder = Path(current_app.instance_path) / "attendance_uploads"
    folder.mkdir(parents=True, exist_ok=True)
    return folder


def normalize_roll(value):
    text = str(value).strip()

    if re.fullmatch(r"\d+\.0", text):
        text = text[:-2]

    return re.sub(r"\s+", "", text).upper()


def get_attendance_data():
    data = session.get("attendance_data")

    if not data:
        return None

    if data.get("teacher_id") != current_user.id:
        return None

    return data


def read_student_file(file_path):
    students = []
    extension = file_path.suffix.lower()

    if extension == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(file_path, data_only=True)

        try:
            sheet = workbook.active

            for row_number, row in enumerate(
                sheet.iter_rows(values_only=True),
                start=1
            ):
                if len(row) < 2:
                    continue

                if row[0] is None or row[1] is None:
                    continue

                roll = normalize_roll(row[0])
                name = str(row[1]).strip()

                if row_number == 1 and roll.lower() in {
                    "roll", "rollno", "roll_no",
                    "rollno.", "rollnumber"
                }:
                    continue

                if roll and name:
                    students.append({
                        "roll": roll,
                        "name": name,
                        "source_row": row_number
                    })

        finally:
            workbook.close()

    elif extension == ".docx":
        from docx import Document

        document = Document(file_path)

        if not document.tables:
            raise ValueError(
                "The Word document must contain a student table."
            )

        table = document.tables[0]

        for row_number, row in enumerate(table.rows, start=1):
            if len(row.cells) < 2:
                continue

            roll = normalize_roll(row.cells[0].text)
            name = row.cells[1].text.strip()

            if row_number == 1 and roll.lower() in {
                "roll", "rollno", "roll_no",
                "rollno.", "rollnumber"
            }:
                continue

            if roll and name:
                students.append({
                    "roll": roll,
                    "name": name,
                    "source_row": row_number
                })

    unique_students = []
    seen = set()

    for student in students:
        if student["roll"] not in seen:
            unique_students.append(student)
            seen.add(student["roll"])

    return unique_students


NUMBER_WORDS = {
    "zero": 0,
    "oh": 0,
    "one": 1,
    "two": 2,
    "three": 3,
    "four": 4,
    "five": 5,
    "six": 6,
    "seven": 7,
    "eight": 8,
    "nine": 9,
    "ten": 10,
    "eleven": 11,
    "twelve": 12,
    "thirteen": 13,
    "fourteen": 14,
    "fifteen": 15,
    "sixteen": 16,
    "seventeen": 17,
    "eighteen": 18,
    "nineteen": 19,
    "twenty": 20,
    "thirty": 30,
    "forty": 40,
    "fifty": 50,
    "sixty": 60,
    "seventy": 70,
    "eighty": 80,
    "ninety": 90
}


def spoken_number_candidates(transcript):
    text = transcript.lower()

    digit_numbers = set(
        re.findall(r"\d+", text)
    )

    tokens = re.findall(r"[a-z]+|\d+", text)
    candidates = set(digit_numbers)

    index = 0

    while index < len(tokens):
        token = tokens[index]

        if token not in NUMBER_WORDS:
            index += 1
            continue

        value = NUMBER_WORDS[token]

        if value >= 20 and value % 10 == 0:
            if index + 1 < len(tokens):
                next_token = tokens[index + 1]

                if next_token in NUMBER_WORDS:
                    next_value = NUMBER_WORDS[next_token]

                    if 1 <= next_value <= 9:
                        value += next_value
                        index += 1

        candidates.add(str(value))
        index += 1

    digit_words = {
        "zero": "0",
        "oh": "0",
        "one": "1",
        "two": "2",
        "three": "3",
        "four": "4",
        "five": "5",
        "six": "6",
        "seven": "7",
        "eight": "8",
        "nine": "9"
    }

    run = []

    for token in tokens + ["END"]:
        if token in digit_words:
            run.append(digit_words[token])
        else:
            if len(run) >= 2:
                candidates.add("".join(run))

            run = []

    return candidates


def find_matches(transcript, students):
    candidates = spoken_number_candidates(transcript)
    matches = []

    for student in students:
        roll = normalize_roll(student["roll"])

        if not roll.isdigit():
            continue

        if roll in candidates:
            matches.append(student)

    return matches


def create_strength_word(file_path, data, present):
    from docx import Document

    document = Document()

    document.add_heading("Attendance Sheet", level=1)
    document.add_paragraph(f'Subject: {data["subject"]}')
    document.add_paragraph(f'Division: {data["division"]}')
    document.add_paragraph(f'Date: {data["date"]}')

    document.add_paragraph(
        f'Total Strength: {len(data["students"])}'
    )

    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"

    headers = table.rows[0].cells
    headers[0].text = "Roll Number"
    headers[1].text = "Student Name"
    headers[2].text = "Attendance"

    for student in data["students"]:
        cells = table.add_row().cells
        roll = normalize_roll(student["roll"])

        cells[0].text = student["roll"]
        cells[1].text = student["name"]
        cells[2].text = "P" if roll in present else "A"

    document.save(file_path)


def add_excel_attendance(file_path, data, present):
    from openpyxl import load_workbook

    workbook = load_workbook(file_path)

    try:
        sheet = workbook.active
        column = sheet.max_column + 1

        sheet.cell(
            row=1,
            column=column,
            value=f'{data["date"]} - {data["subject"]}'
        )

        for student in data["students"]:
            roll = normalize_roll(student["roll"])

            sheet.cell(
                row=student["source_row"],
                column=column,
                value="P" if roll in present else "A"
            )

        workbook.save(file_path)

    finally:
        workbook.close()


def add_word_attendance(file_path, data, present):
    from docx import Document

    document = Document(file_path)

    if not document.tables:
        raise ValueError(
            "The uploaded Word document must contain a student table."
        )

    attendance_table = document.add_table(
        rows=1,
        cols=3
    )

    attendance_table.style = "Table Grid"

    attendance_table.rows[0].cells[0].text = "Roll Number"
    attendance_table.rows[0].cells[1].text = "Student Name"
    attendance_table.rows[0].cells[2].text = "Attendance"

    for student in data["students"]:
        cells = attendance_table.add_row().cells
        roll = normalize_roll(student["roll"])

        cells[0].text = student["roll"]
        cells[1].text = student["name"]
        cells[2].text = "P" if roll in present else "A"

    document.save(file_path)


@attendance_bp.route("/", methods=["GET", "POST"])
@login_required
def attendance_home():
    if request.method == "POST":
        action = request.form.get("action")

        if action in {"strength", "upload"}:
            subject = request.form.get("subject", "").strip()
            division = request.form.get("division", "").strip()
            date_text = request.form.get("lecture_date", "").strip()

            if not subject or not division or not date_text:
                flash(
                    "Enter subject, division, and date.",
                    "error"
                )
                return redirect(
                    url_for("attendance.attendance_home")
                )

            try:
                date_value = datetime.strptime(
                    date_text,
                    "%Y-%m-%d"
                ).date()

            except ValueError:
                flash("Select a valid date.", "error")
                return redirect(
                    url_for("attendance.attendance_home")
                )

            file_id = str(uuid.uuid4())

            if action == "strength":
                try:
                    strength = int(
                        request.form.get("strength", "0")
                    )

                except ValueError:
                    strength = 0

                if not 1 <= strength <= 300:
                    flash(
                        "Enter a strength between 1 and 300.",
                        "error"
                    )
                    return redirect(
                        url_for("attendance.attendance_home")
                    )

                students = [
                    {
                        "roll": str(i),
                        "name": f"Student {i}",
                        "source_row": i + 1
                    }
                    for i in range(1, strength + 1)
                ]

                extension = ".docx"
                mode = "strength"

            else:
                uploaded_file = request.files.get("student_file")

                if not uploaded_file or not uploaded_file.filename:
                    flash("Choose a Word or Excel file.", "error")
                    return redirect(
                        url_for("attendance.attendance_home")
                    )

                extension = Path(
                    secure_filename(uploaded_file.filename)
                ).suffix.lower()

                if extension not in ALLOWED_EXTENSIONS:
                    flash(
                        "Only .docx and .xlsx files are supported.",
                        "error"
                    )
                    return redirect(
                        url_for("attendance.attendance_home")
                    )

                file_path = upload_folder() / f"{file_id}{extension}"
                uploaded_file.save(file_path)

                if file_path.stat().st_size > 5 * 1024 * 1024:
                    file_path.unlink(missing_ok=True)
                    flash("File must be smaller than 5 MB.", "error")
                    return redirect(
                        url_for("attendance.attendance_home")
                    )

                try:
                    students = read_student_file(file_path)

                except Exception:
                    current_app.logger.exception(
                        "Student file reading failed"
                    )
                    file_path.unlink(missing_ok=True)
                    flash("Could not read the student file.", "error")
                    return redirect(
                        url_for("attendance.attendance_home")
                    )

                if not students:
                    file_path.unlink(missing_ok=True)
                    flash(
                        "No students found in the first two columns "
                        "or the first Word table.",
                        "error"
                    )
                    return redirect(
                        url_for("attendance.attendance_home")
                    )

                mode = "upload"

            session["attendance_data"] = {
                "teacher_id": current_user.id,
                "file_id": file_id,
                "extension": extension,
                "mode": mode,
                "subject": subject,
                "division": division,
                "date": date_value.isoformat(),
                "students": students,
                "present": []
            }

            session.modified = True

            return render_template(
                "attendance.html",
                students=students,
                attendance_data=session["attendance_data"]
            )

        if action == "save":
            data = get_attendance_data()

            if not data:
                flash(
                    "Start an attendance session first.",
                    "error"
                )
                return redirect(
                    url_for("attendance.attendance_home")
                )

            present = {
                normalize_roll(roll)
                for roll in data["present"]
            }

            file_path = (
                upload_folder()
                / f'{data["file_id"]}{data["extension"]}'
            )

            try:
                subject = db.session.execute(
                    select(Subject).where(
                        Subject.teacher_id == current_user.id,
                        Subject.name == data["subject"]
                    )
                ).scalar_one_or_none()

                if subject is None:
                    subject = Subject(
                        teacher_id=current_user.id,
                        name=data["subject"]
                    )
                    db.session.add(subject)
                    db.session.flush()

                attendance_session = AttendanceSession(
                    teacher_id=current_user.id,
                    subject_id=subject.id,
                    division=data["division"],
                    session_date=datetime.strptime(
                        data["date"], "%Y-%m-%d"
                    ).date()
                )

                db.session.add(attendance_session)
                db.session.flush()

                for student_data in data["students"]:
                    roll = normalize_roll(student_data["roll"])

                    student = db.session.execute(
                        select(Student).where(
                            Student.teacher_id == current_user.id,
                            Student.roll_number == roll
                        )
                    ).scalar_one_or_none()

                    if student is None:
                        student = Student(
                            teacher_id=current_user.id,
                            roll_number=roll,
                            name=student_data["name"],
                            division=data["division"]
                        )
                        db.session.add(student)
                        db.session.flush()

                    db.session.add(
                        AttendanceRecord(
                            session_id=attendance_session.id,
                            student_id=student.id,
                            status=(
                                "present"
                                if roll in present
                                else "absent"
                            ),
                            marked_by="voice" if data["mode"] == "strength" else "upload"
                        )
                    )

                if data["mode"] == "strength":
                    create_strength_word(file_path, data, present)

                elif data["extension"] == ".xlsx":
                    add_excel_attendance(file_path, data, present)

                else:
                    add_word_attendance(file_path, data, present)

                db.session.commit()

            except Exception:
                db.session.rollback()

                current_app.logger.exception(
                    "Attendance save failed"
                )

                flash(
                    "Could not save attendance. Check the Flask terminal.",
                    "error"
                )

                return redirect(
                    url_for("attendance.attendance_home")
                )

            session.pop("attendance_data", None)

            return send_file(
                file_path,
                as_attachment=True,
                download_name=(
                    f'attendance_{data["date"]}_'
                    f'{secure_filename(data["subject"])}'
                    f'{data["extension"]}'
                )
            )

    data = get_attendance_data()

    return render_template(
        "attendance.html",
        students=data["students"] if data else [],
        attendance_data=data
    )


@attendance_bp.route("/transcribe", methods=["POST"])
@login_required
def transcribe_audio():
    data = get_attendance_data()

    if not data:
        return jsonify({
            "error": "Start an attendance session first."
        }), 400

    audio = request.files.get("audio")

    if not audio:
        return jsonify({
            "error": "No audio received."
        }), 400

    temporary_path = upload_folder() / f"{uuid.uuid4()}.webm"

    try:
        audio.save(temporary_path)

        print(
            f"Received audio: {temporary_path.stat().st_size} bytes",
            flush=True
        )

        if temporary_path.stat().st_size < 1000:
            return jsonify({
                "transcript": "",
                "matches": [],
                "present": data["present"],
                "message": "Waiting for enough speech..."
            })

        from services.speech_service import transcribe_audio_file

        transcript = transcribe_audio_file(str(temporary_path))

        print(
            "Recognized transcript:",
            repr(transcript),
            flush=True
        )

        matches = find_matches(
            transcript,
            data["students"]
        )

        present = {
            normalize_roll(roll)
            for roll in data["present"]
        }

        new_matches = []

        for student in matches:
            roll = normalize_roll(student["roll"])

            if roll not in present:
                data["present"].append(roll)
                present.add(roll)

                new_matches.append({
                    "roll": student["roll"],
                    "name": student["name"]
                })

        session["attendance_data"] = data
        session.modified = True

        message = (
            "Marked present: " + ", ".join(
                item["roll"] for item in new_matches
            )
            if new_matches
            else "No new roll number recognized."
        )

        return jsonify({
            "transcript": transcript,
            "matches": new_matches,
            "present": data["present"],
            "message": message
        })

    except Exception as error:
        print(
            "\n========== TRANSCRIPTION ERROR ==========",
            flush=True
        )

        traceback.print_exc()

        print(
            "Error type:",
            type(error).__name__,
            flush=True
        )

        print(
            "Error details:",
            repr(error),
            flush=True
        )

        print(
            "=========================================\n",
            flush=True
        )

        current_app.logger.exception(
            "Faster-Whisper transcription failed"
        )

        return jsonify({
            "error": str(error),
            "error_type": type(error).__name__
        }), 500

    finally:
        try:
            temporary_path.unlink(missing_ok=True)

        except OSError:
            current_app.logger.exception(
                "Could not delete temporary audio file"
            )